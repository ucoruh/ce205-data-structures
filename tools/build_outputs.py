"""Builds the CEN207 (Data Structures) course outputs locally.

For every week that already has a new-layout lecture note (docs/week-N/cen207-week-N.<lang>.md):
    deck       : slides/.../*.<lang>.md -> docs/.../*-slides.<lang>.html   (Marp, used inside the embedded iframe)
    deck-pdf   : *-slides.<lang>.html   -> docs/.../*-slides.<lang>.pdf    (Chrome print; faster than Marp's own PDF)
    pptx       : slides/.../*.<lang>.md -> docs/.../*-slides.<lang>.pptx   (python-pptx, editable, small)
    note-pdf   : docs/.../*.<lang>.md   -> docs/.../*-notes.<lang>.pdf     (letterhead; from the site HTML via Chrome)
    docx       : docs/.../*.<lang>.md   -> docs/.../*-notes.<lang>.docx    (letterhead; pandoc)
    package    : docs/.../*-materials.<lang>.zip (every output above plus the week's code)
    buttons    : writes/refreshes the per-page download-button block and the embedded deck iframe
    skeleton   : creates a starter deck for a week that has a note but no deck yet (does not touch an existing one)

A week WITHOUT a new-layout note is skipped entirely -- old-layout weeks (docs/week-N-name/...) are left alone
until they are rewritten. Publishing (GitHub Pages) is NOT part of this script; that is separate and explicit.

Usage (from the repository root):
    py -3.12 tools/build_outputs.py all
    py -3.12 tools/build_outputs.py deck deck-pdf buttons --week 3
"""
import argparse
import base64
import functools
import http.server
import os
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile
import threading

ROOT = pathlib.Path(__file__).resolve().parents[1]
DOCS = ROOT / 'docs'
SLIDES = ROOT / 'slides'
THEME_DIR = SLIDES / 'theme'
TOOLS = ROOT / 'tools'

CHROME_CANDIDATES = [
    r'C:\Program Files\Google\Chrome\Application\chrome.exe',
    r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
    '/usr/bin/google-chrome', '/usr/bin/chromium',
]

# ---------------------------------------------------------------- CONFIG (only course-specific values live here)
CONFIG = {
    'course_code': 'CEN207',
    'course_name': {'en': 'Data Structures', 'tr': 'Veri Yapıları'},
    'instructor': {'en': 'Asst. Prof. Dr. Uğur CORUH', 'tr': 'Dr. Öğr. Üyesi Uğur CORUH'},
    'term': {'en': 'Fall 2026-2027', 'tr': '2026-2027 Güz'},
    'institution_name': {'en': 'RECEP TAYYIP ERDOGAN UNIVERSITY', 'tr': 'RECEP TAYYİP ERDOĞAN ÜNİVERSİTESİ'},
    'institution_unit': {
        'en': 'Faculty of Engineering and Architecture · Department of Computer Engineering',
        'tr': 'Mühendislik ve Mimarlık Fakültesi · Bilgisayar Mühendisliği Bölümü',
    },
    'institution_short': {'en': 'RTEU Computer Engineering', 'tr': 'RTEÜ Bilgisayar Mühendisliği'},
    'accent_color': '2E7D32',           # green, no leading '#' (used with RGBColor.from_string and in %-strings)
    'theme_name': 'cen207',
    'file_prefix': 'cen207-week-',
    'primary_lang': 'en',               # course is taught in English; the site's default language is EN
    'secondary_lang': 'tr',             # secondary language; copied from the primary output when not yet written
    'placeholder_marker': 'Bu sayfa hazırlanıyor',   # marks a not-yet-translated TR page (checked, never written by this script)
}
ACCENT = CONFIG['accent_color']
PRIMARY, SECONDARY = CONFIG['primary_lang'], CONFIG['secondary_lang']
PREFIX = CONFIG['file_prefix']
REFERENCE_DOCX = TOOLS / 'reference' / f'{CONFIG["course_code"].lower()}-reference.docx'
HEADER_LOGO = THEME_DIR / 'rteu-logo-small.jpg'          # Marp header + PPTX corner logo
LETTERHEAD_LOGO = DOCS / 'images' / 'rteu_logo.jpg'      # lecture-note PDF/DOCX letterhead

# week: (date, LO codes, English title, Turkish title, English topics, Turkish topics)
# Source: 2026-2027-Guz-CEN207-Data-Structures/izlence/CEN207-2026-2027-Guz-Izlence.{en,tr}.md, section C.
WEEKS = {
    1: ('18.09.2026 (make-up 23.09)', '1, 2, 7', 'Introduction, Big-O, Pointers', 'Giriş, Büyük O, İşaretçiler',
        ['Course plan and communication.',
         'Introduction to linear and non-linear data structures; performance analysis (Big-O).',
         'Pointers and objects for data and variables; memory layout.',
         'Basics of ASN.1 / BER TLV / PER TLV.',
         'Intensive C workshop (toolchain, compile-run-debug).'],
        ['Ders planı ve iletişim.',
         'Doğrusal ve doğrusal olmayan veri yapılarına giriş; performans analizi (Büyük O).',
         'Veri ve değişkenler için işaretçiler ve nesneler; bellek düzeni.',
         'ASN.1 / BER TLV / PER TLV temelleri.',
         'Yoğun C atölyesi (araç zinciri, derle-çalıştır-hata ayıkla).']),
    2: ('25.09.2026', '1, 7', 'Linked Lists, Arrays, Matrices', 'Bağlı Listeler, Diziler, Matrisler',
        ['Linked lists (singly, doubly, circular, XOR) and skip lists.',
         'Arrays (rotation, rearrangement, searching).',
         'Matrices and sparse matrices.'],
        ['Bağlı listeler (tekli, çift, dairesel, XOR) ve atlamalı listeler.',
         'Diziler (döndürme, yeniden düzenleme, arama).',
         'Matrisler ve seyrek matrisler.']),
    3: ('02.10.2026', '1, 7', 'Stacks and Queues', 'Yığınlar ve Kuyruklar',
        ['Stacks (array and linked list, LIFO).',
         'Expressions (infix, postfix, prefix) and conversions.',
         'Queues (standard, circular, deque, multilevel; FIFO).',
         'Tower of Hanoi; recursion (groundwork for DFS).'],
        ['Yığınlar (dizi ve bağlı liste, LIFO).',
         'İfadeler (infix, postfix, prefix) ve dönüşümler.',
         'Kuyruklar (standart, dairesel, çift uçlu, çok seviyeli; FIFO).',
         'Hanoi Kulesi; özyineleme (DFS\'e hazırlık).']),
    4: ('09.10.2026', '1, 4, 7', 'Trees, Heaps, Huffman', 'Ağaçlar, Öbekler, Huffman',
        ['Trees and binary trees; traversals (in-, pre-, post-order).',
         'Heaps (min, max, binary, binomial, Fibonacci, leftist, k-ary) and priority queues; heap sort.',
         'Huffman coding.'],
        ['Ağaçlar ve ikili ağaçlar; dolaşmalar (in-, pre-, post-order).',
         'Yığınlar/heap (min, max, ikili, binom, Fibonacci, leftist, k-ary) ve öncelik kuyrukları; heap sıralama.',
         'Huffman kodlama.']),
    5: ('16.10.2026', '1, 5, 7', 'Graphs and Traversals', 'Çizgeler ve Dolaşmalar',
        ['Graphs: representations (adjacency matrix, incidence matrix, adjacency list).',
         'Traversals (BFS, DFS, iterative deepening, depth-limited, bidirectional).',
         'Topological sorting; water jug problem.'],
        ['Çizgeler: gösterimler (komşuluk matrisi, geliş matrisi, komşuluk listesi).',
         'Dolaşmalar (BFS, DFS, yinelemeli derinleşme, derinlik sınırlı, çift yönlü).',
         'Topolojik sıralama; su kabı problemi.']),
    6: ('23.10.2026', '3, 4, 7', 'Search and Hashing', 'Arama ve Hash',
        ['Searching (linear, binary, interpolation, Fibonacci).',
         'Hashing and hash tables (direct-address tables, hash functions, chaining, open addressing, perfect hashing).',
         'Collision resolution in practice.'],
        ['Arama (doğrusal, ikili, interpolasyon, Fibonacci).',
         'Hashing ve hash tabloları (doğrudan adresli tablolar, hash fonksiyonları, zincirleme, açık adresleme, mükemmel hashing).',
         'Çakışma çözümü uygulamaları.']),
    7: ('30.10.2026', '1–5, 7', 'Midterm Project Demonstrations', 'Ara Proje Gösterimleri',
        ['Midterm project demonstrations (C) and midterm project report.'],
        ['Ara proje gösterimleri (C) ve ara proje raporu.']),
    8: ('31.10–08.11.2026', '1, 2, 4, 5, 7', 'Midterm Exam Week: Quiz-1', 'Ara Sınav Haftası: Quiz-1',
        ['Quiz-1 (weeks 1-6).'], ['Quiz-1 (1-6. haftalar).']),
    9: ('13.11.2026', '3, 5, 7', 'Graph Algorithms', 'Çizge Algoritmaları',
        ['Minimum spanning trees (Prim, Kruskal with disjoint sets).',
         'Shortest paths (Dijkstra, Bellman-Ford); connectivity and SCC; maximum flow.',
         'Cycle detection (Floyd, Brent); backtracking (n-queens, m-coloring, Euler and Hamiltonian paths).'],
        ['Minimum yayılan ağaçlar (Prim, ayrık kümelerle Kruskal).',
         'En kısa yollar (Dijkstra, Bellman-Ford); bağlantılılık ve SCC; maksimum akış.',
         'Döngü tespiti (Floyd, Brent); geri izleme (n-vezir, m-renklendirme, Euler ve Hamilton yolları).']),
    10: ('20.11.2026', '2, 3, 7', 'Sorting', 'Sıralama',
         ['Sorting algorithms and taxonomy (insertion, selection, shell, quick, merge, heap, radix, counting, external sorting).',
          'Comparison of sorting methods.'],
         ['Sıralama algoritmaları ve sınıflandırma (ekleme, seçim, shell, hızlı, birleştirme, heap, radix, sayma, dış sıralama).',
          'Sıralama yöntemlerinin karşılaştırılması.']),
    11: ('27.11.2026', '4, 7', 'Advanced Trees', 'Gelişmiş Ağaçlar',
         ['Binary search trees, AVL, red-black, splay trees.',
          'B-tree family (2-3, 2-3-4, B+, B#); augmenting data structures.',
          'Comparison of search trees.'],
         ['İkili arama ağaçları, AVL, kırmızı-siyah, splay ağaçları.',
          'B-ağacı ailesi (2-3, 2-3-4, B+, B#); veri yapılarının genişletilmesi.',
          'Arama ağaçlarının karşılaştırılması.']),
    12: ('04.12.2026', '1, 3, 4, 7', 'Strings', 'Dizeler',
         ['String structures; search algorithms (brute force, Knuth-Morris-Pratt, Boyer-Moore, Horspool).',
          'LCS and edit distance (Levenshtein, Wagner-Fischer); alignment (Needleman-Wunsch, Smith-Waterman).',
          'Tries and Patricia trees.'],
         ['String yapıları; arama algoritmaları (kaba kuvvet, Knuth-Morris-Pratt, Boyer-Moore, Horspool).',
          'LCS ve düzenleme uzaklığı (Levenshtein, Wagner-Fischer); hizalama (Needleman-Wunsch, Smith-Waterman).',
          'Trie ve Patricia ağaçları.']),
    13: ('11.12.2026', '6, 7', 'File Organisation I', 'Dosya Organizasyonu I',
         ['Sequential files (binary, interpolation, self-organizing search).',
          'Direct files and hashing functions; collision resolution (coalesced hashing, progressive overflow, '
          'double hashing, buckets, Brent\'s method).',
          'Perfect hashing.'],
         ['Sıralı dosyalar (ikili, interpolasyon, kendini düzenleyen arama).',
          'Doğrudan dosyalar ve hash fonksiyonları; çakışma çözümü (birleşik hashing, ilerleyen taşma, '
          'çift hashing, kovalar, Brent yöntemi).',
          'Mükemmel hashing.']),
    14: ('18.12.2026', '3, 4, 6, 7', 'File Organisation II', 'Dosya Organizasyonu II',
         ['Indexed sequential files; secondary key retrieval.',
          'Binary and B-tree structures for files.',
          'Hashing for expandable files (extendible, dynamic, linear hashing); k-d trees and grid files.',
          'External file sorting.'],
         ['İndeksli sıralı dosyalar; ikincil anahtarla erişim.',
          'Dosyalar için ikili ve B-ağacı yapıları.',
          'Genişleyebilen dosyalar için hashing (genişletilebilir, dinamik, doğrusal hashing); k-d ağaçları ve ızgara dosyaları.',
          'Dış dosya sıralama.']),
    15: ('25.12.2026', '1–7', 'Final Project Demonstrations', 'Final Proje Gösterimleri',
         ['Final project demonstrations (Java) and final project report.'],
         ['Final proje gösterimleri (Java) ve final proje raporu.']),
    16: ('04–17.01.2027', '2–7', 'Final Exam Period: Quiz-2', 'Final Sınav Dönemi: Quiz-2',
         ['Quiz-2 (weeks 9-14).'], ['Quiz-2 (9-14. haftalar).']),
}


class Week:
    """File paths for one week's note, deck and generated outputs."""

    # output kind -> filename suffix. mkdocs-static-i18n publishes ".en." at the site root and ".tr." under
    # /tr/ with the SAME (suffix-less) name, so links inside the pages are written without the language suffix.
    KINDS = {
        'deck_html': '-slides.html', 'deck_pdf': '-slides.pdf', 'deck_pptx': '-slides.pptx',
        'note_pdf': '-notes.pdf', 'note_docx': '-notes.docx', 'package': '-materials.zip',
    }

    def __init__(self, number):
        self.number = number
        self.folder = DOCS / f'week-{number}'
        self.name = f'{PREFIX}{number}'
        self.page = {lang: self.folder / f'{self.name}.{lang}.md' for lang in (PRIMARY, SECONDARY)}
        self.deck_source_path = {lang: SLIDES / f'week-{number}' / f'{self.name}.{lang}.md' for lang in (PRIMARY, SECONDARY)}
        w = WEEKS[number]
        self.title = {'en': f'Week {number} — {w[2]}', 'tr': f'Hafta {number} — {w[3]}'}

    @property
    def code_folder(self):
        return ROOT / 'code' / f'week-{self.number:02d}'

    def output_path(self, kind, lang=PRIMARY):
        stem, ext = self.KINDS[kind].rsplit('.', 1)
        return self.folder / f'{self.name}{stem}.{lang}.{ext}'

    def link(self, kind):
        return f'{self.name}{self.KINDS[kind]}'

    def has_note(self):
        return self.page[PRIMARY].exists() or self.page[SECONDARY].exists()

    def deck_source(self, lang):
        """The deck source for that language, or None (the secondary language falls back to a copy
        of the primary language's rendered output when it has no deck of its own yet)."""
        path = self.deck_source_path[lang]
        return path if path.exists() else None

    def has_secondary_deck(self):
        return self.deck_source_path[SECONDARY].exists()

    def has_real_secondary_note(self):
        """True when the secondary-language page is an actual translation, not a placeholder."""
        page = self.page[SECONDARY]
        if not page.exists():
            return False
        text = page.read_text(encoding='utf-8')
        return CONFIG['placeholder_marker'] not in text and len(text.splitlines()) > 60

    def copy_to_secondary(self, kind):
        """Used when the secondary language has no output of its own yet: the primary language's
        output is shown on the secondary-language site too, rather than a broken link."""
        primary_file = self.output_path(kind, PRIMARY)
        if primary_file.exists():
            shutil.copyfile(primary_file, self.output_path(kind, SECONDARY))


def discover_weeks():
    """Only weeks that already have a new-layout note are processed; older-layout weeks are untouched."""
    return [n for n in sorted(WEEKS) if Week(n).has_note()]


def run_cmd(cmd, **kw):
    # A subprocess's output goes through a pipe, Windows then defaults to cp1252, and a line containing a
    # Turkish character (course notes, titles) crashes the caller; force UTF-8 instead.
    kw.setdefault('env', {**os.environ, 'PYTHONUTF8': '1', 'PYTHONIOENCODING': 'utf-8'})
    result = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace', **kw)
    if result.returncode != 0:
        print('   ERROR:', ' '.join(map(str, cmd))[:200])
        print('  ', (result.stderr or result.stdout)[-800:])
    return result.returncode == 0


def chrome_path():
    for candidate in CHROME_CANDIDATES:
        if pathlib.Path(candidate).exists():
            return candidate
    sys.exit('Chrome/Edge not found.')


# ---------------------------------------------------------------- theme and skeleton deck
def build_theme():
    logo = base64.b64encode(HEADER_LOGO.read_bytes()).decode()
    template = (THEME_DIR / f'{CONFIG["theme_name"]}.theme.css.template').read_text(encoding='utf-8')
    (THEME_DIR / f'{CONFIG["theme_name"]}.css').write_text(
        template.replace('{{LOGO}}', f'data:image/jpeg;base64,{logo}'), encoding='utf-8', newline='\n')


FRONT_MATTER = """---
marp: true
theme: {theme}
paginate: true
lang: {lang}
title: "{title}"
author: "{instructor}"
header: "{code} {course_name} · {short_title}"
footer: "{institution_short} · {term}"
---
"""


def skeleton(week):
    for lang in (PRIMARY, SECONDARY):
        source = week.deck_source_path[lang]
        if source.exists():
            continue
        source.parent.mkdir(parents=True, exist_ok=True)
        n = week.number
        date, lo, title_en, title_tr, topics_en, topics_tr = WEEKS[n]
        title = title_en if lang == 'en' else title_tr
        topics = topics_en if lang == 'en' else topics_tr
        bullets = '\n'.join(f'- {t}' for t in topics)
        week_word = 'Week' if lang == 'en' else 'Hafta'
        lo_word = 'Learning outcomes' if lang == 'en' else 'Öğrenme çıktıları'
        prepared = ('The detailed slides for this week will be published together with the lecture notes.'
                    if lang == 'en' else 'Bu haftanın ayrıntılı sunumu ders notuyla birlikte yayımlanacak.')
        text = FRONT_MATTER.format(
            theme=CONFIG['theme_name'], lang=lang, title=f'{CONFIG["course_code"]} {week_word} {n} — {title}',
            instructor=CONFIG['instructor'][lang], code=CONFIG['course_code'],
            course_name=CONFIG['course_name'][lang], short_title=f'{week_word} {n}',
            institution_short=CONFIG['institution_short'][lang], term=CONFIG['term'][lang]) + f"""
<!-- _class: baslik -->
<!-- _paginate: false -->

# {title}

**{CONFIG['course_code']} {CONFIG['course_name'][lang]} — {week_word} {n}**

{CONFIG['instructor'][lang]} · {date}

---

# {'This week' if lang == 'en' else 'Bu hafta'}

{bullets}

**{lo_word}:** LO.{lo}

---

<!-- _class: bolum -->

# {'Slides in preparation' if lang == 'en' else 'Sunum hazırlanıyor'}

{prepared}
"""
        source.write_text(text, encoding='utf-8', newline='\n')
        print('   skeleton deck:', source.relative_to(ROOT))


# ---------------------------------------------------------------- deck HTML / PDF / PPTX
def deck_html(week, lang):
    source = week.deck_source(lang)
    if source is None:
        if lang == PRIMARY:
            print('   no deck source:', week.deck_source_path[PRIMARY].relative_to(ROOT))
        return
    marp = shutil.which('marp')
    if not marp:
        sys.exit('marp not found (npm i -g @marp-team/marp-cli).')
    target = week.output_path('deck_html', lang)
    if run_cmd([marp, str(source), '--config-file', str(SLIDES / 'marp.config.yml'),
                '--theme-set', str(THEME_DIR / f'{CONFIG["theme_name"]}.css'), '--html', '-o', str(target)]):
        # The deck is shown on the page inside an iframe; a link inside it would try to open in that same
        # iframe and most sites (GitHub etc.) refuse to be framed. Every link except an in-deck "#" jump
        # opens in a new tab instead.
        html = target.read_text(encoding='utf-8')
        html = re.sub(r'<a\s(?![^>]*\btarget=)([^>]*\bhref="(?!#)[^"]*"[^>]*)>',
                      r'<a target="_blank" rel="noopener" \1>', html)
        target.write_text(html, encoding='utf-8')
        if lang == PRIMARY and not week.has_secondary_deck():
            week.copy_to_secondary('deck_html')
        print(f'   deck html ({lang}):', target.relative_to(ROOT))


_ANIM_IFRAME_HTML = re.compile(
    r'<iframe\s+class="dsanim"\s+[^>]*\bsrc="([^"?]+?)(?:\?[^"]*)?"[^>]*>\s*</iframe>')


def replace_anim_iframes(html, lang):
    """A slide's <iframe class="dsanim" src="anim/NAME.html?..."> cannot run inside a printed PDF, so it is
    replaced with the animation's last frame (anim/NAME-son.<lang>.png, already produced by tools/dsanim).
    The animation name is always read from the iframe's own src -- never a hard-coded list."""
    def replace(m):
        src = m.group(1)
        p = pathlib.PurePosixPath(src)
        image = p.parent / f'{p.stem}-son.{lang}.png'
        return f'<img class="dsanim" src="{image}" alt="animation">'
    return _ANIM_IFRAME_HTML.sub(replace, html)


def deck_pdf(week, lang):
    if lang == SECONDARY and not week.has_secondary_deck():
        return
    source = week.output_path('deck_html', lang)
    if not source.exists():
        return
    target = week.output_path('deck_pdf', lang)
    chrome = chrome_path()
    # Marp scales code blocks with JavaScript via <pre is="marp-pre">; that element renders EMPTY in
    # Chrome's PDF print. Printing is done from a temporary copy with that element removed (code lines
    # are short here, so scaling is not needed) and with animation iframes swapped for their still image.
    html = source.read_text(encoding='utf-8')
    html = re.sub(r'<pre is="marp-pre"[^>]*>', '<pre>', html)
    html = replace_anim_iframes(html, lang)
    # The print copy is written NEXT TO the html (same folder) so relative image paths (assets/..., anim/...) resolve.
    printable = source.with_name(source.stem + '.print.html')
    try:
        printable.write_text(html, encoding='utf-8')
        ok = run_cmd([chrome, '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                      f'--print-to-pdf={target}', printable.as_uri()])
    finally:
        if printable.exists():
            printable.unlink()
    if ok:
        if lang == PRIMARY and not week.has_secondary_deck():
            week.copy_to_secondary('deck_pdf')
        print(f'   deck pdf ({lang}):', target.relative_to(ROOT))


def deck_pptx(week, lang):
    source = week.deck_source(lang)
    if source is None:
        return
    sys.path.insert(0, str(TOOLS))
    import marp_pptx
    target = week.output_path('deck_pptx', lang)
    count = marp_pptx.convert(str(source), str(target), logo=str(HEADER_LOGO), lang=lang)
    if lang == PRIMARY and not week.has_secondary_deck():
        week.copy_to_secondary('deck_pptx')
    print(f'   deck pptx ({lang}): {target.relative_to(ROOT)} ({count} slides)')


# ---------------------------------------------------------------- per-page download/embed block
MARK_START, MARK_END = '<!-- materials:start -->', '<!-- materials:end -->'


def button_block(week, lang):
    en = lang == 'en'
    labels = {
        'note_pdf': ('Lecture notes (PDF)', 'Ders notu (PDF)', 'material-file-pdf-box'),
        'note_docx': ('Lecture notes (DOCX)', 'Ders notu (DOCX)', 'material-file-word-box'),
        'deck_pdf': ('Slides (PDF)', 'Sunum (PDF)', 'material-presentation'),
        'deck_pptx': ('Slides (PPTX)', 'Sunum (PPTX)', 'material-microsoft-powerpoint'),
        'deck_html': ('Slides (HTML, offline)', 'Sunum (HTML, çevrimdışı)', 'material-language-html5'),
        'package': ('Download all (ZIP)', 'Tümünü indir (ZIP)', 'material-folder-zip'),
    }
    buttons = []
    for kind, (label_en, label_tr, icon) in labels.items():
        if week.output_path(kind, lang).exists():
            link = week.link(kind)
            buttons.append(f'[:{icon}: {label_en if en else label_tr}]({link}){{ .md-button download="{link}" }}')
    has_deck = week.output_path('deck_html', lang).exists()
    lines = [MARK_START, '', '<div class="materials" markdown>', '']
    lines += buttons
    if has_deck:
        full = 'Open slides full screen' if en else 'Sunumu tam ekran aç'
        lines.append(f'[:material-fullscreen: {full}]({week.link("deck_html")})'
                      '{ .md-button .md-button--primary target=_blank }')
    lines += ['', '</div>', '']
    if has_deck:
        title = week.title[lang]
        lines += ['<div class="deck-frame">',
                  f'<iframe src="../{week.link("deck_html")}" title="{title}" loading="lazy" allowfullscreen></iframe>',
                  '</div>', '']
        hint = ('Click inside the slides and use the arrow keys; use the button at the bottom right of the '
                'slides, or the "Open slides full screen" link above, for full screen.') if en else (
                'Sunumun içine tıklayıp ok tuşlarıyla ilerleyin; tam ekran için sunumun sağ altındaki düğmeyi '
                'ya da yukarıdaki "Sunumu tam ekran aç" bağlantısını kullanın.')
        lines += [f'<p class="deck-hint">{hint}</p>', '']
    lines.append(MARK_END)
    return '\n'.join(lines)


def buttons(week):
    for lang, page in week.page.items():
        if not page.exists():
            continue
        text = page.read_text(encoding='utf-8')
        new_block = button_block(week, lang)
        # always (re)place the block right under the title: remove an old copy wherever it is
        text = re.sub(r'\n*' + re.escape(MARK_START) + r'.*?' + re.escape(MARK_END) + r'\n*', '\n\n', text,
                      flags=re.S)
        lines = text.split('\n')
        h1 = next(i for i, s in enumerate(lines) if s.startswith('# '))
        i = h1 + 1
        while i < len(lines) and not lines[i].strip():
            i += 1
        # skip the subtitle line (*CEN207 … *) or an info table directly under the title
        if i < len(lines) and (lines[i].startswith('*') or lines[i].startswith('|')):
            while i < len(lines) and lines[i].strip():
                i += 1
        lines[i:i] = ['', new_block, '']
        text = re.sub(r'\n{3,}', '\n\n', '\n'.join(lines))
        page.write_text(text, encoding='utf-8', newline='\n')
        print(f'   button block: {page.relative_to(ROOT)}')


# ---------------------------------------------------------------- temporary site build and local server
class Site:
    """Builds the site into a temporary folder and serves it from a local (127.0.0.1-only) HTTP server."""

    def __init__(self):
        self.temp_dir = pathlib.Path(tempfile.mkdtemp(prefix='cen207-site-'))
        print(f'[site] temporary build: {self.temp_dir}')
        env = {**os.environ, 'CEN207_GIT': 'false'}
        if not run_cmd([sys.executable, '-m', 'mkdocs', 'build', '--clean', '--site-dir', str(self.temp_dir)],
                        cwd=ROOT, env=env):
            sys.exit('mkdocs build failed.')
        handler = functools.partial(_QuietHandler, directory=str(self.temp_dir))
        self.server = http.server.ThreadingHTTPServer(('127.0.0.1', 0), handler)
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.url = f'http://127.0.0.1:{self.server.server_address[1]}'

    def close(self):
        self.server.shutdown()
        shutil.rmtree(self.temp_dir, ignore_errors=True)


class _QuietHandler(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass


def page_url(week, site_root=None, lang=PRIMARY):
    """The compiled page's path.

    mkdocs-static-i18n puts the DEFAULT language's pages at the site root and every other language under
    its own folder (e.g. tr/). Since the default language can change, the path is picked by checking what
    actually exists rather than being hard-coded.
    """
    tail = f'week-{week.number}/{week.name}/'
    if site_root is None:
        return lang + '/' + tail
    if (site_root / tail / 'index.html').exists() and not (site_root / lang / tail / 'index.html').exists():
        return tail
    return lang + '/' + tail


def logo_data_uri():
    data = base64.b64encode(LETTERHEAD_LOGO.read_bytes()).decode()
    return f'data:image/jpeg;base64,{data}'


# ---------------------------------------------------------------- lecture-note PDF (letterhead)
PRINT_CSS = """
@page {
  size: A4;
  margin: 22mm 16mm 20mm 16mm;
  @top-left   { content: "%(code)s %(course_name)s"; font: 8pt 'Roboto', Arial, sans-serif; color: #5b6770; }
  @top-right  { content: "%(header)s"; font: 8pt 'Roboto', Arial, sans-serif; color: #5b6770; }
  @bottom-left  { content: "%(institution_short)s · %(instructor)s · %(term)s";
                  font: 7.5pt 'Roboto', Arial, sans-serif; color: #5b6770; }
  @bottom-right { content: "Page " counter(page) " / " counter(pages); font: 7.5pt 'Roboto', Arial, sans-serif; color: #5b6770; }
}
@page :first { @top-left { content: none; } @top-right { content: none; } }
html, body { background: #fff !important; }
.md-header, .md-tabs, .md-sidebar, .md-footer, .md-banner, .md-top, .md-source-file, .md-content__button,
.md-consent, .md-dialog, .md-feedback, .md-tags, .materials, .deck-frame, .deck-hint, .headerlink,
.md-announce, aside.md-source-file, .page-metadata, .metadata, .md-tag { display: none !important; }
.md-main__inner, .md-content, .md-content__inner, .md-grid { margin: 0 !important; max-width: none !important; }
.md-content__inner { padding: 0 !important; }
.md-content__inner > :first-child { margin-top: 0 !important; }
.md-typeset { font-size: 9.6pt; line-height: 1.5; }
.md-typeset h1 { font-size: 20pt; color: #%(accent)s; margin: 0.2em 0 0.6em; }
.md-typeset h2 { color: #%(accent)s; border-bottom: 1px solid #C8E6C9; padding-bottom: 2px; break-after: avoid; }
.md-typeset h3, .md-typeset h4 { break-after: avoid; }
.md-typeset pre, .md-typeset table, .md-typeset .admonition, .md-typeset details, .md-typeset .mermaid,
.md-typeset figure, .md-typeset blockquote { break-inside: avoid; }
.md-typeset pre > code { white-space: pre-wrap; word-break: break-word; }
.md-typeset table:not([class]) { font-size: 8.6pt; }
.letterhead { display: flex; align-items: center; gap: 12px; border-bottom: 2.5px solid #%(accent)s;
         padding-bottom: 8px; margin-bottom: 14px; font-family: 'Roboto', Arial, sans-serif; }
.letterhead img { height: 58px; }
.letterhead .middle { flex: 1; line-height: 1.35; }
.letterhead .institution { font-size: 11pt; font-weight: 700; color: #1f2d3a; letter-spacing: .3px; }
.letterhead .unit { font-size: 8.5pt; color: #5b6770; }
.letterhead .course { font-size: 9.5pt; color: #%(accent)s; font-weight: 600; margin-top: 2px; }
.letterhead .right { text-align: right; font-size: 8.5pt; color: #5b6770; line-height: 1.4; }
.letterhead .right b { color: #%(accent)s; font-size: 10pt; }
.md-typeset .tabbed-labels { display: none !important; }
.md-typeset .tabbed-content > .tabbed-block { display: block !important; }
.md-typeset .tab-title { font-weight: 700; color: #%(accent)s; margin: 0.8em 0 0.2em; break-after: avoid; }
"""


def expand_tabs(html):
    """In print, every tab is shown, one after another, each labelled with its own tab title (nested
    tabs are expanded too)."""
    queue = []

    def replace(m):
        if m.group(1) is not None:                                    # tab labels: queue them up
            queue.extend(re.findall(r'<label[^>]*>(.*?)</label>', m.group(1), flags=re.S))
            return m.group(0)
        title = queue.pop(0) if queue else ''
        return m.group(0) + (f'<p class="tab-title">{title}</p>' if title else '')
    return re.sub(r'<div class="tabbed-labels">(.*?)</div>|<div class="tabbed-block">', replace, html, flags=re.S)


LETTERHEAD = """<div class="letterhead">
<img src="%(logo)s" alt="RTEU">
<div class="middle">
<div class="institution">%(institution)s</div>
<div class="unit">%(unit)s</div>
<div class="course">%(code)s %(course_name)s · %(term)s</div>
</div>
<div class="right"><b>%(kind)s</b><br>%(short)s<br>%(date)s</div>
</div>"""


def printable_page(site, week, lang):
    """Writes a print-ready copy of the compiled page (letterhead, expanded tabs) and returns its URL."""
    path = page_url(week, site.temp_dir, lang)
    folder = site.temp_dir / path
    html = (folder / 'index.html').read_text(encoding='utf-8')
    n = week.number
    kind = 'Lecture Note' if lang == 'en' else 'Ders Notu'
    header = f'Week {n} — Lecture Note' if lang == 'en' else f'Hafta {n} — Ders Notu'
    short = f'Week {n}' if lang == 'en' else f'Hafta {n}'
    date = WEEKS[n][0]
    html = re.sub(r'<details(?![^>]*\bopen\b)', '<details open', html)
    html = expand_tabs(html)
    css = PRINT_CSS % {'code': CONFIG['course_code'], 'course_name': CONFIG['course_name'][lang],
                        'header': header, 'institution_short': CONFIG['institution_short'][lang],
                        'instructor': CONFIG['instructor'][lang], 'term': CONFIG['term'][lang], 'accent': ACCENT}
    html = html.replace('</head>', f'<style>{css}</style></head>', 1)
    letterhead = LETTERHEAD % {'logo': logo_data_uri(), 'kind': kind, 'short': short, 'date': date,
                               'institution': CONFIG['institution_name'][lang], 'unit': CONFIG['institution_unit'][lang],
                               'code': CONFIG['course_code'], 'course_name': CONFIG['course_name'][lang], 'term': CONFIG['term'][lang]}
    html, n_sub = re.subn(r'(<article class="?md-content__inner md-typeset"?>)',
                          r'\1' + letterhead.replace('\\', '\\\\'), html, count=1)
    if not n_sub:
        print('   WARNING: article start not found, letterhead not inserted')
    (folder / f'print-{lang}.html').write_text(html, encoding='utf-8')
    return f'{site.url}/{path}print-{lang}.html'


def note_pdf(week, site, lang):
    if lang == PRIMARY and not week.page[PRIMARY].exists():
        return
    if lang == SECONDARY and not week.has_real_secondary_note():
        return
    url = printable_page(site, week, lang)
    target = week.output_path('note_pdf', lang)
    if run_cmd([chrome_path(), '--headless=new', '--disable-gpu', '--no-pdf-header-footer',
                '--virtual-time-budget=20000', '--run-all-compositor-stages-before-draw',
                f'--print-to-pdf={target}', url]):
        if lang == PRIMARY and not week.has_real_secondary_note():
            week.copy_to_secondary('note_pdf')
        print(f'   note pdf ({lang}):', target.relative_to(ROOT))


# ---------------------------------------------------------------- lecture-note DOCX (letterhead, pandoc)
def build_reference_docx():
    """Builds a letterhead reference document from pandoc's own default template (once)."""
    if REFERENCE_DOCX.exists():
        return
    import docx
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    REFERENCE_DOCX.parent.mkdir(parents=True, exist_ok=True)
    raw = REFERENCE_DOCX.with_name('_pandoc_default.docx')
    subprocess.run(['pandoc', '-o', str(raw), '--print-default-data-file', 'reference.docx'], check=True)
    doc = docx.Document(str(raw))

    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin = section.right_margin = Cm(2)
        section.top_margin, section.bottom_margin = Cm(2.6), Cm(2.2)
        section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)
        # Header: logo + institution + course
        header = section.header.paragraphs[0]
        header.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
        header.add_run().add_picture(str(LETTERHEAD_LOGO), height=Cm(1.1))
        r = header.add_run(f'  {CONFIG["institution_name"]["en"]} · {CONFIG["institution_short"]["en"]}')
        r.font.size, r.font.bold, r.font.color.rgb = Pt(8.5), True, RGBColor(0x1F, 0x2D, 0x3A)
        r = header.add_run(f'\t{CONFIG["course_code"]} {CONFIG["course_name"]["en"]}')
        r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(ACCENT)
        border = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '12'), ('w:space', '4'), ('w:color', ACCENT)):
            bottom.set(qn(k), v)
        border.append(bottom)
        header._p.get_or_add_pPr().append(border)
        # Footer: course info + page number
        footer = section.footer.paragraphs[0]
        footer.paragraph_format.tab_stops.add_tab_stop(Cm(17), WD_TAB_ALIGNMENT.RIGHT)
        r = footer.add_run(f'{CONFIG["instructor"]["en"]} · {CONFIG["term"]["en"]}\tPage ')
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x5B, 0x67, 0x70)
        for part in ('begin', 'PAGE', 'end'):
            r = footer.add_run()
            r.font.size = Pt(8)
            if part == 'PAGE':
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = 'PAGE'
            else:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), part)
            r._r.append(el)

    styles = doc.styles

    def set_style(name, size=None, color=None, bold=None, font=None):
        try:
            s = styles[name]
        except KeyError:
            return
        if font:
            s.font.name = font
            s.element.rPr.rFonts.set(qn('w:eastAsia'), font)
        if size:
            s.font.size = Pt(size)
        if color:
            s.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            s.font.bold = bold
    for name in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        set_style(name, 10.5, None, None, 'Calibri')
    set_style('Title', 22, ACCENT, True, 'Calibri')
    set_style('Heading 1', 17, ACCENT, True, 'Calibri')
    set_style('Heading 2', 14, ACCENT, True, 'Calibri')
    set_style('Heading 3', 12, '1F2D3A', True, 'Calibri')
    set_style('Heading 4', 11, '1F2D3A', True, 'Calibri')
    set_style('Source Code', 8.5, None, None, 'Consolas')
    set_style('Verbatim Char', 9, 'A1123F', None, 'Consolas')
    # Admonitions (Block Text): left rule + soft background
    try:
        bt = styles['Block Text']
        ppr = bt.element.get_or_add_pPr()
        border = OxmlElement('w:pBdr')
        left = OxmlElement('w:left')
        for k, v in (('w:val', 'single'), ('w:sz', '18'), ('w:space', '8'), ('w:color', ACCENT)):
            left.set(qn(k), v)
        border.append(left)
        ppr.append(border)
        shade = OxmlElement('w:shd')
        for k, v in (('w:val', 'clear'), ('w:color', 'auto'), ('w:fill', 'EAF6EC')):
            shade.set(qn(k), v)
        ppr.append(shade)
        bt.font.italic = False
    except KeyError:
        pass
    # Tables: thin borders
    try:
        table_style = styles['Table']
        tblpr = table_style.element.find(qn('w:tblPr'))
        if tblpr is None:
            tblpr = OxmlElement('w:tblPr')
            table_style.element.append(tblpr)
        borders = OxmlElement('w:tblBorders')
        for k in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            e = OxmlElement(f'w:{k}')
            for a, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '0'), ('w:color', 'C8E6C9')):
                e.set(qn(a), v)
            borders.append(e)
        tblpr.append(borders)
    except KeyError:
        pass
    doc.save(str(REFERENCE_DOCX))
    raw.unlink()
    print('   reference docx:', REFERENCE_DOCX.relative_to(ROOT))


def mermaid_png(site, code, target):
    """Renders Mermaid code on a temporary page served by the local site and saves it as a (cropped) PNG."""
    from PIL import Image, ImageChops
    folder = site.temp_dir / '_mermaid'
    folder.mkdir(exist_ok=True)
    page = folder / (target.stem + '.html')
    safe = code.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
    page.write_text(
        '<!doctype html><html><head><meta charset="utf-8"><style>body{margin:0;background:#fff}'
        '#d{display:inline-block;padding:12px}</style></head><body><div id="d"><pre class="mermaid">'
        + safe + '</pre></div><script type="module">import m from '
        '"https://cdn.jsdelivr.net/npm/mermaid@11/dist/mermaid.esm.min.mjs";'
        'm.initialize({startOnLoad:true,theme:"default",fontFamily:"Arial",'
        'flowchart:{useMaxWidth:false,htmlLabels:true}});</script></body></html>',
        encoding='utf-8')
    run_cmd([chrome_path(), '--headless=new', '--disable-gpu', '--hide-scrollbars', '--window-size=1400,1800',
              '--force-device-scale-factor=2', '--virtual-time-budget=10000', f'--screenshot={target}',
              f'{site.url}/_mermaid/{page.name}'])
    if not target.exists():
        return None
    img = Image.open(target).convert('RGB')
    box = ImageChops.difference(img, Image.new('RGB', img.size, (255, 255, 255))).getbbox()
    if box:
        img = img.crop((max(box[0] - 10, 0), max(box[1] - 10, 0), box[2] + 10, box[3] + 10))
    img.save(target)
    return img.size


def docx_preprocess(text, site, temp_dir, page_folder=None, lang=None):
    """Rewrites MkDocs-specific Markdown into plain Markdown that pandoc understands."""
    text = re.sub(r'\A---\n.*?\n---\n', '', text, flags=re.S)                        # front matter
    text = re.sub(re.escape(MARK_START) + r'.*?' + re.escape(MARK_END), '', text, flags=re.S)  # download/embed block
    text = re.sub(r'\{ *\.md-button[^}]*\}', '', text)

    # Data-structure animations: the iframe player cannot run in Word, so it is dropped, and its
    # <div class="dsanim-baski" markdown> fallback (a plain Markdown image, already written by the note)
    # is unwrapped so the ordinary image-handling step below picks it up.
    text = re.sub(r'<iframe\s+class="dsanim"[^>]*>\s*</iframe>', '', text)
    text = re.sub(r'<div class="dsanim-baski"[^>]*>\s*(.*?)\s*</div>', r'\1', text, flags=re.S)

    # Local images (assets/...): a PNG sibling is used instead of SVG (pandoc/Word cannot reliably embed
    # SVG). This runs BEFORE the Mermaid step so Mermaid's own PNG links are not touched here.
    def local_image(m):
        from PIL import Image as _Im
        alt, path = m.group(1), m.group(2)
        if path.startswith(('http://', 'https://', 'data:')) or page_folder is None:
            return m.group(0)
        source = pathlib.Path(page_folder) / path
        png = source.with_suffix('.png')
        if not png.exists():
            png = source
        if not png.exists() and lang:
            # the site (mkdocs-static-i18n, suffix mode) serves NAME.<lang>.png as NAME.png; on disk use the suffixed file
            png = source.with_name(f'{source.stem}.{lang}{source.suffix}')
        if not png.exists():
            return '*(image on the course page)*'
        target = temp_dir / png.name
        if not target.exists():
            shutil.copy2(png, target)
        try:
            with _Im.open(target) as im:
                gw = im.size[0]
            width = min(16.0, gw / 2 * 2.54 / 96)
        except Exception:
            width = 15.0
        return f'![{alt}]({target.name}){{width={width:.1f}cm}}'
    text = re.sub(r'!\[([^\]]*)\]\(([^)]+)\)', local_image, text)

    # Mermaid -> PNG
    counter = [0]

    def mermaid(m):
        counter[0] += 1
        png = temp_dir / f'diagram-{counter[0]}.png'
        size = mermaid_png(site, m.group(1), png)
        if not size:
            return '*(diagram on the web page)*'
        width = min(16.0, size[0] / 2 * 2.54 / 96)
        return f'![]({png.name}){{width={width:.1f}cm}}'
    text = re.sub(r'```mermaid\n(.*?)\n```', mermaid, text, flags=re.S)

    # Code-block titles: ```c title="x"  ->  **x** + ```c
    text = re.sub(r'^(\s*)```(\w+) +title="([^"]+)"', r'\1**\3**\n\n\1```\2', text, flags=re.M)

    # Admonitions (!!!, ???, ???+) -> a titled blockquote; tabs (=== "Title") -> a bold title + content.
    # Nested ones are expanded too.
    def expand_boxes(lines):
        out, i = [], 0
        admonition = re.compile(r'^(\s*)(!!!|\?\?\?\+?)\s+(\w+)(?:\s+"([^"]*)")?\s*$')
        tab = re.compile(r'^(\s*)===\+?\s+"([^"]*)"\s*$')
        while i < len(lines):
            m = admonition.match(lines[i])
            t = None if m else tab.match(lines[i])
            if not m and not t:
                out.append(lines[i])
                i += 1
                continue
            indent = len((m or t).group(1))
            body, i = [], i + 1
            while i < len(lines) and (not lines[i].strip() or
                                       len(lines[i]) - len(lines[i].lstrip()) >= indent + 4):
                body.append(lines[i][indent + 4:] if lines[i].strip() else '')
                i += 1
            while body and not body[-1]:
                body.pop()
            body = expand_boxes(body)
            if t:
                prefix = ' ' * indent
                out += [prefix + f'**{t.group(2)}**', '']
                out += [(prefix + b) if b else '' for b in body]
                out.append('')
                continue
            title = m.group(4) or m.group(3).capitalize()
            prefix = ' ' * indent + '> '
            out.append(prefix + f'**{title}**')
            out.append(prefix.rstrip())
            out += [(prefix + b) if b else prefix.rstrip() for b in body]
            out.append('')
        return out
    text = '\n'.join(expand_boxes(text.split('\n')))
    text = re.sub(r'<p class="deck-hint">.*?</p>', '', text, flags=re.S)
    return text


def note_docx(week, site, lang):
    if lang == PRIMARY and not week.page[PRIMARY].exists():
        return
    if lang == SECONDARY and not week.has_real_secondary_note():
        return
    source_page = week.page[lang]
    build_reference_docx()
    temp_dir = pathlib.Path(tempfile.mkdtemp(prefix='cen207-docx-'))
    try:
        md = docx_preprocess(source_page.read_text(encoding='utf-8'), site, temp_dir, page_folder=source_page.parent, lang=lang)
        source = temp_dir / 'note.md'
        source.write_text(md, encoding='utf-8')
        target = week.output_path('note_docx', lang)
        title = week.title[lang]
        if run_cmd(['pandoc', str(source), '-f', 'markdown+pipe_tables+grid_tables+fenced_code_attributes'
                    '+link_attributes+raw_html-implicit_figures', '-t', 'docx', '--reference-doc', str(REFERENCE_DOCX),
                    '--resource-path', str(temp_dir), '--metadata', f'title-meta={title}', '-o', str(target)]):
            if lang == PRIMARY and not week.has_real_secondary_note():
                week.copy_to_secondary('note_docx')
            print(f'   note docx ({lang}):', target.relative_to(ROOT))
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


# ---------------------------------------------------------------- offline package (ZIP)
def package(week, lang):
    """Bundles every material for the week, plus its demo code, into one ZIP (build output excluded)."""
    import zipfile
    target = week.output_path('package', lang)
    root_name = week.name
    # Matches .gitignore: build/run output and anything a setup script downloads never goes into the ZIP.
    skip = {'bin', 'dokum', 'build', '__pycache__', 'cikti', 'lib'}
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for kind in ('deck_html', 'deck_pdf', 'deck_pptx', 'note_pdf', 'note_docx'):
            file = week.output_path(kind, lang)
            if file.exists():
                z.write(file, f'{root_name}/{week.link(kind)}')
        code = week.code_folder
        if code and code.exists():
            # The week's demos plus the shared build infrastructure (root CMake, scripts, cmake/, common/):
            # the ZIP's code/ folder builds on its own once extracted.
            code_root = code.parent
            common = [code_root / n for n in ('CMakeLists.txt', 'CMakePresets.json', 'build.ps1', 'build.sh',
                                               'README.md', 'README.en.md', '.gitattributes')]
            for folder in (code_root / 'cmake', code_root / 'common', code):
                common += sorted(folder.rglob('*'))
            for path in common:
                parts = set(path.relative_to(code_root).parts)
                if (path.is_file() and not (parts & skip) and path.name.lower() != 'desktop.ini'
                        and path.suffix.lower() != '.jar'):
                    z.write(path, f'{root_name}/code/{path.relative_to(code_root).as_posix()}')
    print(f'   package: {target.relative_to(ROOT)} ({target.stat().st_size // 1024} KB)')


# ---------------------------------------------------------------- main flow
STEPS = ['skeleton', 'deck', 'deck-pdf', 'pptx', 'note-pdf', 'docx', 'package', 'buttons']


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('steps', nargs='+', choices=STEPS + ['all'])
    ap.add_argument('--week', help='comma-separated week numbers (default: every week with a new-layout note)')
    args = ap.parse_args()
    steps = STEPS if 'all' in args.steps else args.steps
    numbers = [int(n) for n in args.week.split(',')] if args.week else discover_weeks()
    build_theme()
    weeks = [Week(n) for n in numbers]
    site = None
    for step in STEPS:
        if step not in steps:
            continue
        if step in ('note-pdf', 'docx'):
            site = Site()
        print(f'== {step}')
        for week in weeks:
            if step == 'skeleton':
                skeleton(week)
            elif step == 'deck':
                deck_html(week, PRIMARY)
                deck_html(week, SECONDARY)
            elif step == 'deck-pdf':
                deck_pdf(week, PRIMARY)
                deck_pdf(week, SECONDARY)
            elif step == 'pptx':
                deck_pptx(week, PRIMARY)
                deck_pptx(week, SECONDARY)
            elif step == 'note-pdf':
                note_pdf(week, site, PRIMARY)
                note_pdf(week, site, SECONDARY)
            elif step == 'docx':
                note_docx(week, site, PRIMARY)
                note_docx(week, site, SECONDARY)
            elif step == 'package':
                package(week, PRIMARY)
                package(week, SECONDARY)
            elif step == 'buttons':
                buttons(week)
        if site:
            site.close()
            site = None


if __name__ == '__main__':
    main()
